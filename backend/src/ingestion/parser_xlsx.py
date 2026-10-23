"""Bounded XLSX parsing for the first worksheet of authorized registry datasets."""
from contextlib import contextmanager
from io import BytesIO
from zipfile import BadZipFile, ZipFile
from openpyxl import load_workbook
from .storage import download_spooled_file

MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_XLSX_ZIP_ENTRIES = 2000
MAX_XLSX_COLUMNS = 500

def _validate_xlsx(file_obj):
    file_obj.seek(0)
    try:
        with ZipFile(file_obj) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ZIP_ENTRIES:
                raise ValueError("XLSX contains too many archive entries")
            expanded = sum(entry.file_size for entry in entries)
            if expanded > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError("XLSX exceeds the uncompressed size limit")
            names = set(archive.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("File is not a valid XLSX workbook")
            if any(entry.file_size > MAX_XLSX_UNCOMPRESSED_BYTES for entry in entries):
                raise ValueError("XLSX archive entry exceeds the size limit")
    except (BadZipFile, OSError) as exc:
        raise ValueError("XLSX file is malformed") from exc
    finally:
        file_obj.seek(0)

def iter_xlsx_file(file_obj, max_records: int):
    if max_records < 1:
        raise ValueError("max_records must be positive")
    _validate_xlsx(file_obj)
    workbook = None
    try:
        workbook = load_workbook(file_obj, read_only=True, data_only=True)
        if not workbook.worksheets:
            raise ValueError("XLSX workbook has no worksheets")
        sheet = workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if not header_row or len(header_row) > MAX_XLSX_COLUMNS:
            raise ValueError("XLSX has no usable header row or too many columns")
        headers = [str(value).strip() if value is not None else "" for value in header_row]
        if not any(headers):
            raise ValueError("XLSX has no usable header row")
        for row_number, values in enumerate(rows, start=2):
            if row_number - 1 > max_records:
                raise ValueError(f"XLSX exceeds configured limit of {max_records} records")
            record = {}
            used = set()
            for index, value in enumerate(values[:len(headers)]):
                header = headers[index]
                if not header:
                    continue
                key = header
                if key in used:
                    key = f"{header}_{index + 1}"
                used.add(key)
                record[key] = "" if value is None else str(value).strip()
            if any(value for value in record.values()):
                yield row_number, record
    finally:
        if workbook is not None:
            workbook.close()

def iter_xlsx_records(data: bytes, max_records: int):
    with BytesIO(data) as file_obj:
        yield from iter_xlsx_file(file_obj, max_records)

def iter_xlsx_records_from_storage(key: str, max_records: int):
    with download_spooled_file(key) as file_obj:
        yield from iter_xlsx_file(file_obj, max_records)
